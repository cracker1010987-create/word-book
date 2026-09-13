# NGSL/TSL 원본 엑셀을 읽어 앱에서 쓸 단어장(data/word_bank.json)을 만드는 스크립트
import csv

from openpyxl import load_workbook


# 원본 엑셀의 첫 시트에서 제목 줄을 빼고, A열 단어와 B열 영어 뜻을 순서대로 읽어온다
def read_source_words(xlsx_path: str) -> list[dict]:
    sheet = load_workbook(xlsx_path, read_only=True).active
    words = []
    for row in list(sheet.iter_rows(values_only=True))[1:]:
        word = str(row[0] or "").strip()
        if not word:
            continue
        definition = str(row[1] or "").strip() if len(row) > 1 else ""
        words.append({"word": word, "definition": definition})
    return words


# 통계 CSV 파일의 모든 줄을 읽는다. utf-8로 안 읽히면(TSL 파일) latin-1로 다시 읽는다
def read_csv_rows(csv_path: str) -> list[list[str]]:
    try:
        with open(csv_path, encoding="utf-8", newline="") as f:
            return list(csv.reader(f))
    except UnicodeDecodeError:
        with open(csv_path, encoding="latin-1", newline="") as f:
            return list(csv.reader(f))


# 통계 CSV에서 제목 줄을 빼고 1열 단어와 2열 순위를 읽는다 (빈도 지수 열은 #N/A가 있어서 안 쓴다)
def read_source_ranks(csv_path: str) -> list[dict]:
    ranks = []
    for row in read_csv_rows(csv_path)[1:]:
        if len(row) < 2 or not row[0].strip() or not row[1].strip().isdigit():
            continue
        ranks.append({"word": row[0].strip(), "rank": int(row[1])})
    return ranks
